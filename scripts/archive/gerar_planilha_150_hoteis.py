import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import json

json_path = r"C:\Users\Augusto\Desktop\AH3\scripts\sorted_150_hotels.json"
with open(json_path, "r", encoding="utf-8") as f:
    hotels_data = json.load(f)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Top 150 Hotéis Floripa"

# Styling
fill_header = PatternFill(start_color="1B2420", end_color="1B2420", fill_type="solid")
font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

fill_top10 = PatternFill(start_color="FFF8E1", end_color="FFF8E1", fill_type="solid")
fill_accent_green = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
fill_accent_yellow = PatternFill(start_color="FFFDE7", end_color="FFFDE7", fill_type="solid")

font_body = Font(name="Calibri", size=10)
font_bold = Font(name="Calibri", size=10, bold=True)
font_link = Font(name="Calibri", size=10, color="0D3B66", underline="single")

thin_border = Border(
    left=Side(style='thin', color='E0E0E0'),
    right=Side(style='thin', color='E0E0E0'),
    top=Side(style='thin', color='E0E0E0'),
    bottom=Side(style='thin', color='E0E0E0')
)

headers = [
    "Ranking", "Hotel / Empreendimento", "Unidade / Especificação", "Categoria",
    "Região / Bairro", "Nº de Quartos (Capacidade)", "Diária Média (R$)", "Site Oficial / Link",
    "Identidade Visual", "Plaquinha 10x10", "Proposta em PDF", "Status da Venda", "Valor Fechado"
]

ws.append(headers)

for col_num in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col_num)
    cell.fill = fill_header
    cell.font = font_header
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws.row_dimensions[1].height = 28

for h in hotels_data:
    rank = h["ranking"]
    nome = h["nome"]
    unidade = h["unidade"]
    cat = h["categoria"]
    reg = h["regiao"]
    quartos = h["quartos"]
    diaria = h["diaria_media"]
    site = h["site"]
    
    status_id = "Concluído" if rank <= 100 else "Pendente"
    status_plaq = "Concluído" if rank <= 31 else "Pendente"
    status_pdf = "Pendente"
    status_venda = "Não Iniciado"
    valor_venda = 0.00
    
    row_values = [
        rank, nome, unidade, cat, reg, quartos, diaria, site,
        status_id, status_plaq, status_pdf, status_venda, valor_venda
    ]
    ws.append(row_values)
    current_row = ws.max_row
    ws.row_dimensions[current_row].height = 20
    
    for col_idx in range(1, len(row_values) + 1):
        cell = ws.cell(row=current_row, column=col_idx)
        cell.border = thin_border
        cell.font = font_body
        
        if col_idx == 1:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.font = font_bold
            if rank <= 10:
                cell.fill = fill_top10
        elif col_idx == 2:
            cell.font = font_bold
            cell.alignment = Alignment(horizontal="left", vertical="center")
        elif col_idx in [3, 4, 5]:
            cell.alignment = Alignment(horizontal="left", vertical="center")
        elif col_idx == 6:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.font = font_bold
            cell.number_format = '#,##0'
        elif col_idx == 7:
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.font = font_bold
            cell.number_format = '"R$ "#,##0.00'
        elif col_idx == 8:
            cell.font = font_link
            cell.alignment = Alignment(horizontal="left", vertical="center")
        elif col_idx in [9, 10, 11]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if cell.value == "Concluído":
                cell.fill = fill_accent_green
            elif cell.value == "Em Andamento":
                cell.fill = fill_accent_yellow
        elif col_idx == 12:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif col_idx == 13:
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.number_format = '"R$ "#,##0.00'

col_widths = {
    1: 10, 2: 36, 3: 32, 4: 28, 5: 24, 6: 18, 7: 18, 8: 40,
    9: 18, 10: 18, 11: 18, 12: 18, 13: 16
}

for col_idx, width in col_widths.items():
    ws.column_dimensions[get_column_letter(col_idx)].width = width

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:M{len(hotels_data) + 1}"

out_path = r"C:\Users\Augusto\Desktop\AH3\Top_150_Hoteis_Florianopolis_Ranking.xlsx"
wb.save(out_path)
print("Planilha Top 150 salva em:", out_path)
